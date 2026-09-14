from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class ParsedSection:
    text: str
    metadata: dict = field(default_factory=dict)


def parse_file(file_path: str, file_type: str) -> list[ParsedSection]:
    path = Path(file_path)
    if file_type == "pdf":
        return _parse_pdf(path)
    if file_type == "docx":
        return _parse_docx(path)
    if file_type == "pptx":
        return _parse_pptx(path)
    if file_type == "xlsx":
        return _parse_xlsx(path)
    if file_type == "csv":
        return _parse_csv(path)
    if file_type == "json":
        return _parse_json(path)
    if file_type == "xml":
        return _parse_xml(path)
    if file_type in {"md", "markdown", "txt", "html", "htm"}:
        return [ParsedSection(text=path.read_text(encoding="utf-8", errors="ignore"), metadata={"filename": path.name})]
    if file_type in _CODE_EXTENSIONS:
        return [ParsedSection(text=path.read_text(encoding="utf-8", errors="ignore"), metadata={"filename": path.name, "file_type": file_type})]
    return [ParsedSection(text=path.read_text(encoding="utf-8", errors="ignore"), metadata={"filename": path.name, "file_type": file_type})]


_CODE_EXTENSIONS = {
    "py", "js", "ts", "tsx", "jsx", "java", "go", "rs", "c", "cpp", "h", "hpp",
    "cs", "rb", "php", "sh", "sql", "yml", "yaml", "toml", "json", "xml",
}


def _parse_pdf(path: Path) -> list[ParsedSection]:
    import fitz

    sections: list[ParsedSection] = []
    with fitz.open(path) as doc:
        for page_index, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text.strip():
                sections.append(
                    ParsedSection(
                        text=text,
                        metadata={"page": page_index, "filename": path.name, "has_text_layer": True},
                    )
                )
    if not sections:
        sections.append(
            ParsedSection(
                text="",
                metadata={"filename": path.name, "has_text_layer": False, "ocr_required": True},
            )
        )
    return sections


def _parse_docx(path: Path) -> list[ParsedSection]:
    from docx import Document

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return [ParsedSection(text="\n".join(paragraphs), metadata={"filename": path.name})]


def _parse_pptx(path: Path) -> list[ParsedSection]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    sections: list[ParsedSection] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
        if parts:
            sections.append(ParsedSection(text="\n".join(parts), metadata={"page": slide_index, "filename": path.name}))
    if not sections:
        sections.append(ParsedSection(text="", metadata={"filename": path.name}))
    return sections


def _parse_xlsx(path: Path) -> list[ParsedSection]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[ParsedSection] = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                rows.append(" | ".join(values))
        sections.append(ParsedSection(text="\n".join(rows), metadata={"sheet": sheet.title, "filename": path.name}))
    return sections


def _parse_csv(path: Path) -> list[ParsedSection]:
    import csv

    with path.open(encoding="utf-8", errors="ignore", newline="") as file:
        rows = [" | ".join(row) for row in csv.reader(file) if row]
    return [ParsedSection(text="\n".join(rows), metadata={"filename": path.name})]


def _parse_json(path: Path) -> list[ParsedSection]:
    import json

    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    lines: list[str] = []
    _flatten_json(data, "$", lines)
    return [ParsedSection(text="\n".join(lines), metadata={"filename": path.name, "format": "json"})]


def _flatten_json(value, prefix: str, lines: list[str]) -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            _flatten_json(child, f"{prefix}.{key}", lines)
    elif isinstance(value, list):
        for index, child in enumerate(value):
            _flatten_json(child, f"{prefix}[{index}]", lines)
    else:
        lines.append(f"{prefix}: {value}")


def _parse_xml(path: Path) -> list[ParsedSection]:
    from lxml import etree

    tree = etree.parse(str(path))
    lines: list[str] = []
    _flatten_xml(tree.getroot(), lines)
    return [ParsedSection(text="\n".join(lines), metadata={"filename": path.name, "format": "xml"})]


def _flatten_xml(element, lines: list[str], path: str = "") -> None:
    current = f"{path}/{element.tag}"
    text = (element.text or "").strip()
    if text:
        lines.append(f"{current}: {text}")
    for child in element:
        _flatten_xml(child, lines, current)